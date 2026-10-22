#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
actualizar_semividas.py — replace the half-lives of the inventory workbooks with NUBASE2020

For every nuclide row of the "AZ6 RZ1 …" sheets:
  * the name (column B) is rebuilt from its ZZAAA code (column C), which is the
    identifier given by Kairos Power; 16 lanthanides had the wrong element name;
  * the NUBASE2020 half-life is written in new columns (text and seconds), and
    column M (seconds) and N (λ) now read from them, so the whole chain
    (atoms → activity → attenuation by barrier) recalculates;
  * stable nuclides get λ = 0, and so zero activity.
The "Hoja1" sheet (Anexo C of the memoria) is rebuilt with formulas pointing to
the recalculated rows: radioactive nuclides only, sorted by core activity.

Usage:
  python datos/nubase2020/actualizar_semividas.py datos/inventario/Calculosfueradelacontencion.xlsx [datos/inventario/ECC_Calculations.xlsx …]
Writes <name>_NUBASE2020.xlsx next to each input and cambios_semividas.csv here.
"""

from __future__ import annotations

import csv, math, sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from nubase import YEAR_S, half_life

SYM = ("n H He Li Be B C N O F Ne Na Mg Al Si P S Cl Ar K Ca Sc Ti V Cr Mn Fe Co Ni Cu Zn Ga Ge As Se "
       "Br Kr Rb Sr Y Zr Nb Mo Tc Ru Rh Pd Ag Cd In Sn Sb Te I Xe Cs Ba La Ce Pr Nd Pm Sm Eu Gd Tb Dy Ho "
       "Er Tm Yb Lu Hf Ta W Re Os Ir Pt Au Hg Tl Pb Bi Po At Rn Fr Ra Ac Th Pa U Np Pu Am Cm Bk Cf").split()
# Double-beta emitters with T½ above this are treated as stable (no activity of interest)
PRACTICALLY_STABLE_S = 1e15 * YEAR_S
HERE = Path(__file__).resolve().parent
COL_TXT, COL_S, COL_OLD = 58, 59, 60          # BF, BG, BH — free columns after BD
SRC = "NUBASE2020 (Kondev et al., Chin. Phys. C 45, 030001, 2021)"


def name_of(zzaaa: int) -> str:
    return f"{SYM[zzaaa // 1000]}-{zzaaa % 1000}"


# Chemical group of each element for the attenuation factors (column D of the workbook)
GROUP = {"Xe": "Gases nobles", "Kr": "Gases nobles",
         "I": "I, Br, Se, Te", "Br": "I, Br, Se, Te", "Se": "I, Br, Se, Te", "Te": "I, Br, Se, Te",
         "Cs": "Cs, Rb", "Rb": "Cs, Rb", "Sr": "Sr, Ba, Eu", "Ba": "Sr, Ba, Eu", "Eu": "Sr, Ba, Eu",
         "Ag": "Ag, Pd", "Pd": "Ag, Pd", "Mo": "Mo, Ru, Rh, Tc", "Ru": "Mo, Ru, Rh, Tc",
         "Rh": "Mo, Ru, Rh, Tc", "Tc": "Mo, Ru, Rh, Tc",
         "La": "Isot tierras raras", "Ce": "Isot tierras raras", "Pr": "Isot tierras raras",
         "Nd": "Isot tierras raras", "Pm": "Isot tierras raras", "Sm": "Isot tierras raras",
         "Gd": "Isot tierras raras", "Tb": "Isot tierras raras", "Dy": "Isot tierras raras",
         "Y": "Isot elem trans", "Zr": "Isot elem trans", "Nb": "Isot elem trans",
         "U": "Pu, Actinidos", "Np": "Pu, Actinidos", "Pu": "Pu, Actinidos", "Am": "Pu, Actinidos",
         "Cm": "Pu, Actinidos"}
AF_VALUES = range(22, 36)        # V..AI: attenuation factors (50 % / 95 %) of each barrier
AF_FORMULAS = range(36, 57)      # AJ..BD: source term after each barrier


def fix_groups(ws, radioactive_rows):
    """Give every radioactive nuclide the attenuation factors of its own chemical group.

    Rows renamed from ZZAAA could carry the group of the wrong element, and nuclides
    previously treated as stable had no factors at all. Values and formulas are copied
    from a reference row of the right group. Returns {row: note}.
    """
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter as L
    ref = {}
    for r in range(4, ws.max_row + 1):
        g = ws.cell(r, 4).value
        if g and ws.cell(r, 50).value is not None and g not in ref and ws.cell(r, COL_OLD).value is None:
            ref[g] = r
    notes = {}
    for r, name in radioactive_rows:
        want = GROUP.get(name.split("-")[0])
        have = ws.cell(r, 4).value
        missing = ws.cell(r, 50).value is None
        if want is None or want not in ref or (have == want and not missing):
            continue
        src = ref[want]
        ws.cell(r, 4, want)
        for c in AF_VALUES:
            ws.cell(r, c, ws.cell(src, c).value)
        for c in AF_FORMULAS:
            f = ws.cell(src, c).value
            if isinstance(f, str) and f.startswith("="):
                f = Translator(f, origin=f"{L(c)}{src}").translate_formula(f"{L(c)}{r}")
            ws.cell(r, c, f)
        notes[r] = (f"grupo «{have}» → «{want}»" if have and have != want else f"grupo «{want}» asignado") + \
                   f" (factores copiados de {ws.cell(src, 2).value})"
    return notes


AF_FLIBE = (32, 33)               # AF, AG: FLiBe attenuation factor at 50 % and 95 %


def fix_noble_gases(ws):
    """Noble gases are not retained by the salt: AF_FLiBe = 1.

    KP-TR-012-NP-A, Appendix B, Table B-4 gives a cumulative FLiBe release fraction
    of 1.0 for noble gases and tritium. The workbook used 100 for every group, which
    divided the noble gas release by 100 (non-conservative). Returns {row: note}.
    """
    notes = {}
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 4).value != "Gases nobles":
            continue
        old = [ws.cell(r, c).value for c in AF_FLIBE]
        if all(v == 1 for v in old):
            continue
        for c in AF_FLIBE:
            ws.cell(r, c, 1)
        notes[r] = f"AF FLiBe {old[0]}/{old[1]} → 1 (gas noble, KP-TR-012 Tabla B-4)"
    return notes


def update_sheet(ws, wb_values, report, book):
    ws_v = wb_values[ws.title]
    ws.cell(3, COL_TXT, "T½ NUBASE2020").font = Font(bold=True)
    ws.cell(3, COL_S, "T½ NUBASE2020 (s)").font = Font(bold=True)
    ws.cell(3, COL_OLD, "Nombre anterior (corregido)").font = Font(bold=True)
    ws.cell(3, COL_TXT).comment = Comment(f"Fuente: {SRC}. 1 y = 365,2422 d.", "TFM")
    for c in (6, 7, 8):  # old half-life columns, kept only for traceability
        h = ws.cell(3, c).value or ""
        if "sustituida" not in h:
            ws.cell(3, c, f"{h.strip()} — sustituida por NUBASE2020")
    rows = []
    for r in range(4, ws.max_row + 1):
        zz = ws.cell(r, 3).value
        if not isinstance(zz, (int, float)) or ws.cell(r, 13).value is None:
            continue
        zz = int(zz)
        old_name = (ws.cell(r, 2).value or "").strip()
        new_name = name_of(zz)
        old_s = ws_v.cell(r, 13).value
        old_A = ws_v.cell(r, 19).value
        try:
            T, txt = half_life(zz)
        except KeyError:
            T, txt = None, "no figura en NUBASE2020"
        if T is not None and T > PRACTICALLY_STABLE_S:
            txt += " (estable a efectos prácticos)"
            T = None
        ws.cell(r, 2, new_name)
        if new_name != old_name:
            ws.cell(r, COL_OLD, old_name)
        ws.cell(r, COL_TXT, txt)
        ws.cell(r, COL_S, T)
        ws.cell(r, 13, f"=IF(ISNUMBER(BG{r}),BG{r},\"\")")
        ws.cell(r, 14, f"=IF(ISNUMBER(M{r}),LN(2)/M{r},0)")
        nb = ws_v.cell(r, 5).value
        vol = ws_v.cell(r, 16).value
        new_A = nb * 1e24 * vol * math.log(2) / T if T else 0.0
        rows.append((r, new_name, new_A))
        report.append(dict(libro=book, hoja=ws.title, fila=r, zzaaa=zz, nombre=new_name,
                           nombre_anterior=old_name if old_name != new_name else "",
                           semivida_anterior_s=old_s if isinstance(old_s, (int, float)) else "",
                           semivida_nubase=txt, semivida_nubase_s=T if T else "",
                           actividad_anterior_Bq=old_A if isinstance(old_A, (int, float)) else "",
                           actividad_nueva_Bq=new_A, grupo=""))
    if "ACCIDENT" in ws.title:          # only this layout has the per-barrier chain in V..BD
        notes = fix_groups(ws, [(r, n) for r, n, a in rows if a > 0])
        for r, note in fix_noble_gases(ws).items():
            notes[r] = (notes[r] + "; " if r in notes else "") + note
        for item in report:
            if item["hoja"] == ws.title and item["libro"] == book and item["fila"] in notes:
                item["grupo"] = notes[item["fila"]]
        ws.cell(3, COL_OLD + 1, "Grupo químico corregido").font = Font(bold=True)
        for r, note in notes.items():
            ws.cell(r, COL_OLD + 1, note)
    return rows


def rebuild_annex(wb, sheet_name, rows):
    """Hoja1 = Anexo C: radioactive nuclides, formulas to the 95 % attenuation chain."""
    if "Hoja1" not in wb.sheetnames:
        return
    h = wb["Hoja1"]
    head = next((r for r in range(1, h.max_row + 1) if h.cell(r, 5).value == "Core Activity"), None)
    if head is None:
        return
    for r in range(head + 1, max(h.max_row, head + 200) + 1):
        for c in range(4, 12):
            h.cell(r, c).value = None
    src = "'%s'" % sheet_name
    radio = sorted((x for x in rows if x[2] > 0), key=lambda x: -x[2])
    for i, (r, name, _) in enumerate(radio, start=head + 1):
        h.cell(i, 4, name)
        for j, col in enumerate(("AX", "AY", "AZ", "BA", "BB", "BC", "BD")):
            h.cell(i, 5 + j, f"={src}!{col}{r}").number_format = "0.00000E+00"
    h.cell(head + len(radio) + 2, 4,
           f"Semividas: {SRC}. Solo nucleidos radiactivos, ordenados por actividad en el núcleo.")
    return len(radio)


def main(paths):
    report = []
    for p in map(Path, paths):
        wb = load_workbook(p)
        wb_values = load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            if ws.title.startswith("AZ6"):
                rows = update_sheet(ws, wb_values, report, p.name)
                if "ACCIDENT" in ws.title:
                    n = rebuild_annex(wb, ws.title, rows)
                    if n:
                        print(f"{p.name}: Hoja1 rebuilt with {n} radioactive nuclides")
        out = p.with_name(p.stem + "_NUBASE2020.xlsx")
        wb.save(out)
        print("written", out)
    with (HERE / "cambios_semividas.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(report[0]))
        w.writeheader()
        w.writerows(report)
    print("report:", HERE / "cambios_semividas.csv", f"({len(report)} rows)")


if __name__ == "__main__":
    main(sys.argv[1:])
