#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_nif.py — RADTRAD inventory (.nif) from the attenuated source term

Source: the "Reactor Building" column (activity outside the reactor building, 95 %
attenuation factors) of datos/inventario/Calculosfueradelacontencion_NUBASE2020.xlsx.
That is, the inventory already includes the release from the fuel and the attenuation
of every barrier; RADTRAD must then release it as is (see the notes at the end).

For each radioactive nuclide it writes the 12-token line read by PIPELINE.generate_nif:
  number  name  group  T½ (s)  A  inventory (Ci/MWth)  daughter1 BR1  daughter2 BR2  daughter3 BR3
  * names rebuilt from ZZAAA, half-lives from NUBASE2020;
  * RADTRAD release groups 1–9 (noble gases, iodine, caesium, tellurium, strontium,
    barium, ruthenium, cerium, lanthanum), as in the template inventory files;
  * inventory = activity outside the building (Ci) / plant power (MWth), because RADTRAD
    multiplies the .nif by the plant power level of the .psf (PIPELINE sets it to 320 MWth);
  * daughters from the NUBASE2020 decay modes, kept only if they are in the inventory.

Usage:
  python generar_nif.py [workbook] [--potencia 320] [--salida Variables_NIF_terminofuente.txt]
"""

from __future__ import annotations

import argparse, re, sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "datos" / "nubase2020"))
from nubase import NUBASE_FILE, half_life  # noqa: E402

DEFAULT_BOOK = ROOT / "datos" / "inventario" / "Calculosfueradelacontencion_NUBASE2020.xlsx"
SHEET = "AZ6 RZ1 (ACCIDENT)"
COL_NAME, COL_ZZ, COL_BUILDING = 2, 3, 56        # B, C, BD (95 % chain, after the building)
BQ_PER_CI = 3.7e10

SYM = ("n H He Li Be B C N O F Ne Na Mg Al Si P S Cl Ar K Ca Sc Ti V Cr Mn Fe Co Ni Cu Zn Ga Ge As Se "
       "Br Kr Rb Sr Y Zr Nb Mo Tc Ru Rh Pd Ag Cd In Sn Sb Te I Xe Cs Ba La Ce Pr Nd Pm Sm Eu Gd Tb Dy Ho "
       "Er Tm Yb Lu Hf Ta W Re Os Ir Pt Au Hg Tl Pb Bi Po At Rn Fr Ra Ac Th Pa U Np Pu Am Cm Bk Cf").split()

# RADTRAD release groups (same numbering as the template inventory files)
GROUP = {**{e: 1 for e in ("Xe", "Kr")}, **{e: 2 for e in ("I", "Br")}, **{e: 3 for e in ("Cs", "Rb")},
         **{e: 4 for e in ("Te", "Sb", "Se")}, "Sr": 5, "Ba": 6,
         **{e: 7 for e in ("Ru", "Rh", "Pd", "Mo", "Tc", "Co", "Ag")},
         **{e: 8 for e in ("Ce", "Pu", "Np", "U")},
         **{e: 9 for e in ("La", "Y", "Zr", "Nb", "Pr", "Nd", "Pm", "Sm", "Eu", "Gd", "Am", "Cm")}}


def name_of(zzaaa: int) -> str:
    return f"{SYM[zzaaa // 1000]}-{zzaaa % 1000}"


def decay_modes() -> dict:
    """{zzaaa: [(daughter_zzaaa, fraction), …]} for ground states, from the NUBASE2020 BR field."""
    out = {}
    for ln in NUBASE_FILE.read_text(encoding="latin-1").splitlines():
        if ln.startswith("#") or len(ln) < 120 or ln[7] != "0":
            continue
        a, z = int(ln[0:3]), int(ln[4:7])
        modes = []
        for item in ln[119:].split(";"):
            m = re.match(r"\s*(B-|B\+|EC|e\+|A)\s*[=~<>]+\s*([\d.]+)", item)
            if not m:
                continue
            kind, pct = m.group(1), float(m.group(2))
            dz, da = {"B-": (1, 0), "B+": (-1, 0), "EC": (-1, 0), "e+": (-1, 0), "A": (-2, -4)}[kind]
            modes.append(((z + dz) * 1000 + a + da, pct / 100.0))
        # EC and B+ are often given both as "EC+B+" totals; merge equal daughters
        merged = {}
        for d, f in modes:
            merged[d] = max(merged.get(d, 0.0), f)
        out[z * 1000 + a] = sorted(merged.items(), key=lambda x: -x[1])
    return out


def build(book: Path, power: float):
    ws = load_workbook(book, data_only=True)[SHEET]
    modes = decay_modes()
    rows = []
    for r in range(4, ws.max_row + 1):
        zz = ws.cell(r, COL_ZZ).value
        act = ws.cell(r, COL_BUILDING).value
        if not isinstance(zz, (int, float)) or not isinstance(act, (int, float)) or act <= 0:
            continue
        zz = int(zz)
        T, _ = half_life(zz)
        if T is None:
            continue
        rows.append((zz, act))
    present = {zz for zz, _ in rows}
    lines, total_ci = [], 0.0
    for k, (zz, act) in enumerate(sorted(rows, key=lambda x: (x[0] // 1000, x[0] % 1000)), 1):
        name = name_of(zz)
        T, _ = half_life(zz)
        ci = act / BQ_PER_CI
        total_ci += ci
        daughters = [(name_of(d), f) for d, f in modes.get(zz, []) if d in present][:3]
        daughters += [("none", 0.0)] * (3 - len(daughters))
        tok = [f"{k:03d}", name, str(GROUP[name.split('-')[0]]), f"{T:.6E}", str(zz % 1000), f"{ci / power:.6E}"]
        for d, f in daughters:
            tok += [d, f"{f:.5E}"]
        lines.append(" ".join(tok))
    return lines, total_ci


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("libro", nargs="?", type=Path, default=DEFAULT_BOOK)
    ap.add_argument("--potencia", type=float, default=320.0, help="plant power level of the .psf [MWth]")
    ap.add_argument("--salida", type=Path, default=HERE / "Variables_NIF_terminofuente.txt")
    a = ap.parse_args()
    lines, total = build(a.libro, a.potencia)
    a.salida.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"{a.salida}: {len(lines)} nuclides, {total:.4E} Ci outside the building "
          f"({total / a.potencia:.4E} Ci/MWth at {a.potencia:g} MWth)")
    # also write the .nif itself, with the generator used by PIPELINE.py
    sys.path.insert(0, str(HERE))
    from PIPELINE import generate_nif
    nif = a.salida.with_name("kpfhr_terminofuente.nif")
    generate_nif(a.salida, nif, title="KP-FHR attenuated source term (Ci/MWth)")
    print(f"{nif}: copy it to Defaults\\ and choose it when PIPELINE.py asks for the .nif")


if __name__ == "__main__":
    main()
